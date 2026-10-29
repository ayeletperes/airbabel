"""Ingest the KIARVA supplement (Corcoran et al. 2026, Immunity).

KIARVA is the **Karolinska Institutet Adaptive Immune Receptor Gene Variant Atlas** —
the name of the *set*, not of an author. It was previously credited here to
"Kirsanov et al." with DOI ``10.1016/j.immuni.2026.01.012``; neither is right (that
DOI resolves to an unrelated macrophage paper). The atlas is introduced by name in
Corcoran et al., *Immunity* 2026;59(4):1107-1122.e5. See
``allele_server/sources/kiarva.toml`` for the verified citation.

Reads the per-segment "validations summary" sheets of the ``mmc3.xlsx`` and
ingests just the **allele names + sequences** (db_name, sequence, sequence_AA).
The per-subject genotype / population-count sheets are repertoire data and are
deliberately skipped (out of this server's scope).
"""

from __future__ import annotations

import re
from pathlib import Path

from sqlalchemy.orm import Session

from ..curie import build as curie_build
from ..naming import locus_of, scheme_label, segment_of
from ..sequence import SequenceError
from ..species import normalize_species
from . import loader

# The KIARVA paper. This string is the KIARVA row's `sources.version`, so changing it
# creates a new row rather than updating the old one — a full reload clears the stale
# one; `allele-sources check` reports it if you only re-run this pass.
_DOI = "10.1016/j.immuni.2026.01.026"
# The workbook carries a validations-summary sheet per segment and the manifest
# already declares segments = ["V", "D", "J"]; only V was ever read, so KIARVA's
# curated D and J were sitting in a file we open on every run.
_SHEETS = (
    "IGHV validations summary",
    "IGHD validations summary",
    "IGHJ validations summary",
)

# Flanking-padded variants (db_name carries a `_F<n>` tag): their sequence is the
# coding allele plus genomic flanks, but the CODING is identical to the base
# allele. Dropped to keep coding-only.
_FLANKING_RE = re.compile(r"_F\d+")


def ingest_ki_xlsx(
    session: Session, path: str | Path, *, species_key: str = "human",
    sheets=_SHEETS,
) -> dict[str, int]:
    """Ingest KI validation sheets (names + sequences + provenance). Returns counts."""
    import openpyxl

    tag = normalize_species(species_key)
    source = loader.get_or_create_source(
        session, name="KIARVA", version=_DOI, curie_prefix="KIARVA",
        note="KIARVA — Karolinska Institutet Adaptive Immune Receptor Gene Variant Atlas",
    )
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    stats = {"coding": 0, "attributions": 0, "skipped": 0, "skipped_flanking": 0}
    for sheet in sheets:
        if sheet not in wb.sheetnames:
            continue
        it = wb[sheet].iter_rows(values_only=True)
        header = list(next(it))
        idx = {h: i for i, h in enumerate(header)}
        ic, iseq, iaa = idx.get("db_name"), idx.get("sequence"), idx.get("sequence_AA")
        if ic is None or iseq is None:
            continue
        for row in it:
            name = str(row[ic] or "").strip()
            seq = str(row[iseq] or "").strip()
            if _FLANKING_RE.search(name):  # coding == base allele; skip the padded form
                stats["skipped_flanking"] += 1
                continue
            if not name or not seq or "*" not in name or len(name) > 64:
                continue
            segment = segment_of(name)
            if segment not in ("V", "D", "J"):
                stats["skipped"] += 1
                continue
            try:
                coding = loader.upsert_coding(
                    session, species=tag, segment=segment, locus=locus_of(name), nt=seq
                )
            except SequenceError:
                stats["skipped"] += 1
                continue
            stats["coding"] += 1
            loader.add_name(session, coding, source, name, scheme_label(name))
            # Provenance. This was missing entirely, so KIARVA contributed sequences
            # and names while holding zero sequence_source rows -- and the 209 alleles
            # only KIARVA publishes ended up with no source at all, against the
            # "attribution on every record" rule. db_name is KIARVA's own per-allele
            # identifier and is already capped at 64 chars above.
            loader.add_attribution(
                session, coding, source,
                original_id=curie_build(source.curie_prefix or "KIARVA", name),
                original_sequence=seq,
            )
            stats["attributions"] += 1
            # AA is V-only; D and J carry no protein.
            aa = str(row[iaa]).strip() if iaa is not None and row[iaa] else None
            if aa and segment == "V":
                loader.set_v_aa(session, coding, aa_source=aa, source=source)
    wb.close()
    return stats
